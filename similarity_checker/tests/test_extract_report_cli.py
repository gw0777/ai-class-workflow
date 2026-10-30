import json
import os
import tempfile

from similarity_checker import cli
from similarity_checker.compare import compare_texts
from similarity_checker.config import CheckConfig
from similarity_checker.extract import extract, iter_files
from similarity_checker.report import to_html, to_json, to_markdown

A = "협동학습은 학생들이 소집단을 이루어 공동의 목표를 달성하기 위해 서로 돕는 수업 모형이다.\n새롭고 독창적인 문장이 하나 더 있다 여기에 길게."
B = "협동학습은 학생들이 소집단을 이루어 공동의 목표를 달성하기 위해 서로 돕는 수업 모형이다."


def _make_docs(d):
    import docx
    from pptx import Presentation

    p_txt = os.path.join(d, "a.txt")
    open(p_txt, "w", encoding="utf-8").write(A)
    p_md = os.path.join(d, "b.md")
    open(p_md, "w", encoding="utf-8").write("# 제목\n\n**" + B + "**")
    doc = docx.Document(); doc.add_paragraph(B); t = doc.add_table(rows=1, cols=2); t.cell(0, 0).text = "표셀"; doc.save(os.path.join(d, "c.docx"))
    prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = "슬라이드 제목"; s.placeholders[1].text = B
    prs.save(os.path.join(d, "d.pptx"))
    open(os.path.join(d, "e.html"), "w", encoding="utf-8").write(f"<html><body><h1>HTML</h1><p>{B}</p><script>x=1</script></body></html>")
    open(os.path.join(d, "ignored.bin"), "wb").write(b"\x00")


def test_extract_formats_and_iter_files():
    with tempfile.TemporaryDirectory() as d:
        _make_docs(d)
        files = iter_files([d])
        assert [os.path.basename(f) for f in files] == ["a.txt", "b.md", "c.docx", "d.pptx", "e.html"]
        for f in files:
            doc = extract(f)
            assert "협동학습은" in doc.text or "독창적인" in doc.text, f
        assert extract(os.path.join(d, "b.md")).title == "제목"
        assert "표셀" in extract(os.path.join(d, "c.docx")).text
        assert "슬라이드 제목" in extract(os.path.join(d, "d.pptx")).text
        assert "x=1" not in extract(os.path.join(d, "e.html")).text


def test_reports_render():
    res = compare_texts(A, B, CheckConfig(), label_a="A", label_b="B")
    res.references = [{"raw": "김철수(2020). 제목.", "status": "미확인", "note": "없음", "doi": None, "url": None, "matched_title": ""}]
    res.existence = {"title": "T", "hits": [{"provider": "crossref", "title": "T", "score": 1.0, "url": "https://x"}], "errors": ["openalex: x"], "verdict": "v"}
    html = to_html(res)
    assert "<mark>" in html and "참고문헌 실존 검증" in html and "기존 문헌 유무" in html
    md = to_markdown(res)
    assert "표절률" in md and "## 의심 문장" in md
    data = json.loads(to_json(res))
    assert data["suspect_sentences"] == 1 and data["sources"][0]["source"] == "B"


def test_cli_end_to_end(capsys):
    with tempfile.TemporaryDirectory() as d:
        _make_docs(d)
        db = os.path.join(d, "db", "corpus.db")
        assert cli.main(["--db", db, "index", "add", os.path.join(d, "c.docx"), os.path.join(d, "d.pptx")]) == 0
        assert cli.main(["--db", db, "index", "list"]) == 0
        out = capsys.readouterr().out
        assert "총 2건" in out
        out_dir = os.path.join(d, "out")
        rc = cli.main(["--db", db, "check", os.path.join(d, "a.txt"), "--out-dir", out_dir, "--refs", "--fail-above", "10"])
        assert rc == 2  # 표절률 기준 초과 → CI 실패 코드
        assert sorted(os.listdir(out_dir)) == ["a_report.html", "a_report.json", "a_report.md"]
        data = json.load(open(os.path.join(out_dir, "a_report.json"), encoding="utf-8"))
        assert data["suspect_sentences"] == 1 and data["references"] == []
        assert cli.main(["compare", os.path.join(d, "a.txt"), os.path.join(d, "b.md"), "-q"]) == 0
        assert cli.main(["--db", db, "index", "remove", "1", "2"]) == 0
        assert cli.main(["--db", db, "index", "stats"]) == 0
        assert cli.main(["extract", os.path.join(d, "b.md"), "--sentences"]) == 0
