"""문서 → 텍스트 추출 (PDF / DOCX / PPTX / HWPX / TXT / MD / HTML)."""
from __future__ import annotations

import logging
import os
import re
import shutil
import subprocess
import zipfile
from dataclasses import dataclass, field

from .textproc import fix_spaced_hangul, normalize_text

SUPPORTED_EXTS = {".pdf", ".docx", ".pptx", ".hwpx", ".txt", ".md", ".markdown", ".html", ".htm"}


@dataclass
class ExtractedDoc:
    path: str
    text: str
    title: str = ""
    kind: str = ""
    method: str = ""
    pages: int = 0
    warnings: list[str] = field(default_factory=list)

    @property
    def needs_ocr(self) -> bool:
        return self.kind == "pdf" and len(self.text.strip()) < 20 * max(self.pages, 1)


def guess_title(text: str, fallback: str) -> str:
    for ln in text.split("\n"):
        ln = ln.strip()
        if 2 <= len(ln) <= 120 and not re.fullmatch(r"[\d\W_]+", ln):
            return ln
    return fallback


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _pdf_pdfplumber(path: str) -> tuple[str, int]:
    import pdfplumber  # type: ignore

    logging.getLogger("pdfminer").setLevel(logging.ERROR)
    parts = []
    with pdfplumber.open(path) as pdf:
        n = len(pdf.pages)
        for page in pdf.pages:
            parts.append(page.extract_text() or "")
    return "\n\n".join(parts), n


def _pdf_pypdf(path: str) -> tuple[str, int]:
    from pypdf import PdfReader  # type: ignore

    r = PdfReader(path)
    parts = [(p.extract_text() or "") for p in r.pages]
    return fix_spaced_hangul("\n\n".join(parts)), len(r.pages)


def _pdf_pdftotext(path: str) -> tuple[str, int]:
    exe = shutil.which("pdftotext")
    if not exe:
        raise RuntimeError("pdftotext 없음")
    out = subprocess.run([exe, "-enc", "UTF-8", path, "-"], capture_output=True, text=True, timeout=300)
    if out.returncode != 0:
        raise RuntimeError(out.stderr.strip())
    text = out.stdout
    return fix_spaced_hangul(text), text.count("\f") + 1


PDF_ENGINES = {"pdftotext": _pdf_pdftotext, "pdfplumber": _pdf_pdfplumber, "pypdf": _pdf_pypdf}
# 기본 순서: pdftotext(poppler, 가장 빠름) → pdfplumber(한글 띄어쓰기 품질 우수) → pypdf.
# 환경변수 SIMCHECK_PDF_ENGINE=pdfplumber 처럼 지정하면 해당 추출기를 우선 사용한다.


def _engine_order() -> list[str]:
    order = ["pdftotext", "pdfplumber", "pypdf"]
    pref = os.environ.get("SIMCHECK_PDF_ENGINE", "").strip().lower()
    if pref in PDF_ENGINES:
        order.remove(pref)
        order.insert(0, pref)
    return order


def extract_pdf(path: str) -> ExtractedDoc:
    doc = ExtractedDoc(path=path, text="", kind="pdf")
    best_text, best_pages, best_method = "", 0, ""
    for name in _engine_order():
        fn = PDF_ENGINES[name]
        try:
            text, pages = fn(path)
        except Exception as e:  # 라이브러리 미설치/파싱 실패 → 다음 추출기
            doc.warnings.append(f"{name} 실패: {e}")
            continue
        if len(text.strip()) > len(best_text.strip()):
            best_text, best_pages, best_method = text, pages, name
        if len(text.strip()) >= 80 * max(pages, 1):  # 충분히 추출됨
            break
    doc.text, doc.pages, doc.method = best_text, best_pages, best_method
    if doc.needs_ocr:
        doc.warnings.append("텍스트가 거의 없습니다(스캔본 추정). OCR 후 TXT로 저장해 다시 검사하세요.")
    return doc


# --------------------------------------------------------------------------- #
# Office / HWPX
# --------------------------------------------------------------------------- #
def extract_docx(path: str) -> ExtractedDoc:
    import docx  # type: ignore

    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return ExtractedDoc(path=path, text="\n".join(parts), kind="docx", method="python-docx")


def extract_pptx(path: str) -> ExtractedDoc:
    from pptx import Presentation  # type: ignore

    prs = Presentation(path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            parts.append(slide.notes_slide.notes_text_frame.text)
        parts.append("")
    return ExtractedDoc(path=path, text="\n".join(parts), kind="pptx", method="python-pptx", pages=len(prs.slides))


_TAG_RE = re.compile(r"<[^>]+>")


def extract_hwpx(path: str) -> ExtractedDoc:
    """HWPX(zip+XML)는 외부 라이브러리 없이 문단 텍스트만 뽑는다."""
    parts = []
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n))
        for n in names:
            xml = z.read(n).decode("utf-8", errors="ignore")
            for para in re.findall(r"<hp:p\b.*?</hp:p>", xml, flags=re.S):
                runs = re.findall(r"<hp:t[^>]*>(.*?)</hp:t>", para, flags=re.S)
                if runs:
                    parts.append(_TAG_RE.sub("", "".join(runs)))
    text = "\n".join(parts).replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    return ExtractedDoc(path=path, text=text, kind="hwpx", method="zip-xml", pages=len(names))


def extract_html(path: str) -> ExtractedDoc:
    raw = open(path, encoding="utf-8", errors="ignore").read()
    raw = re.sub(r"<(script|style)\b.*?</\1>", " ", raw, flags=re.S | re.I)
    raw = re.sub(r"</(p|div|li|h\d|tr|br)\s*>", "\n", raw, flags=re.I)
    text = _TAG_RE.sub(" ", raw)
    text = text.replace("&nbsp;", " ").replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    return ExtractedDoc(path=path, text=text, kind="html", method="regex")


def extract_plain(path: str) -> ExtractedDoc:
    for enc in ("utf-8", "utf-8-sig", "cp949", "euc-kr"):
        try:
            text = open(path, encoding=enc).read()
            break
        except UnicodeDecodeError:
            continue
    else:
        text = open(path, encoding="utf-8", errors="ignore").read()
    kind = "md" if path.lower().endswith((".md", ".markdown")) else "txt"
    if kind == "md":
        text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)        # 헤딩 기호
        text = re.sub(r"[*_`]{1,3}", "", text)                     # 강조
        text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)           # 이미지
        text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)       # 링크
        text = re.sub(r"^\|?[\s:|-]+\|?$", "", text, flags=re.M)   # 표 구분선
    return ExtractedDoc(path=path, text=text, kind=kind, method="read")


# --------------------------------------------------------------------------- #
# 진입점
# --------------------------------------------------------------------------- #
def extract(path: str) -> ExtractedDoc:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        doc = extract_pdf(path)
    elif ext == ".docx":
        doc = extract_docx(path)
    elif ext == ".pptx":
        doc = extract_pptx(path)
    elif ext == ".hwpx":
        doc = extract_hwpx(path)
    elif ext in (".html", ".htm"):
        doc = extract_html(path)
    elif ext == ".hwp":
        raise ValueError("HWP(구형 바이너리)는 지원하지 않습니다. 한글에서 HWPX/DOCX/PDF로 저장 후 사용하세요.")
    elif ext in (".txt", ".md", ".markdown", ""):
        doc = extract_plain(path)
    else:
        raise ValueError(f"지원하지 않는 형식: {ext} (지원: {', '.join(sorted(SUPPORTED_EXTS))})")
    doc.text = normalize_text(doc.text)
    doc.title = guess_title(doc.text, os.path.basename(path))
    return doc


def iter_files(paths: list[str]) -> list[str]:
    """파일/디렉터리 목록을 지원 형식 파일 목록으로 펼친다."""
    out: list[str] = []
    for p in paths:
        if os.path.isdir(p):
            for root, _dirs, files in os.walk(p):
                for f in sorted(files):
                    if os.path.splitext(f)[1].lower() in SUPPORTED_EXTS:
                        out.append(os.path.join(root, f))
        elif os.path.isfile(p):
            out.append(p)
    return out
