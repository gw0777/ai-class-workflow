import os
from typing import Dict, List, Any
import PyPDF2
from pptx import Presentation


class PresentationParser:
    """Parse PDF and PowerPoint files to extract content"""

    @staticmethod
    def parse_pdf(file_path: str) -> Dict[str, Any]:
        """
        Parse PDF file and extract text content from each page

        Args:
            file_path: Path to PDF file

        Returns:
            Dictionary containing slides information
        """
        slides = []

        try:
            # Use PyPDF2 for PDF text extraction
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                total_pages = len(pdf_reader.pages)

                for page_num in range(total_pages):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text() or ""

                    slide_info = {
                        'slide_number': page_num + 1,
                        'content': text.strip(),
                        'has_tables': False,  # PyPDF2 doesn't detect tables
                        'tables_count': 0,
                        'word_count': len(text.split()),
                        'char_count': len(text)
                    }

                    slides.append(slide_info)

            return {
                'success': True,
                'file_type': 'pdf',
                'total_slides': total_pages,
                'slides': slides,
                'total_words': sum(s['word_count'] for s in slides),
                'total_chars': sum(s['char_count'] for s in slides)
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e),
                'file_type': 'pdf'
            }

    @staticmethod
    def parse_pptx(file_path: str) -> Dict[str, Any]:
        """
        Parse PowerPoint file and extract text content from each slide

        Args:
            file_path: Path to PPTX file

        Returns:
            Dictionary containing slides information
        """
        slides = []

        try:
            prs = Presentation(file_path)
            total_slides = len(prs.slides)

            for slide_num, slide in enumerate(prs.slides, 1):
                text_content = []

                # Extract text from all shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_content.append(shape.text)

                # Join all text
                full_text = "\n".join(text_content)

                # Check for notes
                notes_text = ""
                if slide.has_notes_slide:
                    notes_slide = slide.notes_slide
                    if notes_slide.notes_text_frame:
                        notes_text = notes_slide.notes_text_frame.text

                slide_info = {
                    'slide_number': slide_num,
                    'content': full_text.strip(),
                    'notes': notes_text.strip(),
                    'has_notes': bool(notes_text.strip()),
                    'shapes_count': len(slide.shapes),
                    'word_count': len(full_text.split()),
                    'char_count': len(full_text)
                }

                slides.append(slide_info)

            return {
                'success': True,
                'file_type': 'pptx',
                'total_slides': total_slides,
                'slides': slides,
                'total_words': sum(s['word_count'] for s in slides),
                'total_chars': sum(s['char_count'] for s in slides)
            }

        except Exception as e:
            return {
                'success': False,
                'error': str(e),
                'file_type': 'pptx'
            }

    @staticmethod
    def parse_file(file_path: str) -> Dict[str, Any]:
        """
        Parse file based on extension

        Args:
            file_path: Path to file

        Returns:
            Parsed content dictionary
        """
        ext = os.path.splitext(file_path)[1].lower()

        if ext == '.pdf':
            return PresentationParser.parse_pdf(file_path)
        elif ext in ['.ppt', '.pptx']:
            return PresentationParser.parse_pptx(file_path)
        else:
            return {
                'success': False,
                'error': f'Unsupported file type: {ext}'
            }

    @staticmethod
    def get_summary(parsed_data: Dict[str, Any]) -> Dict[str, Any]:
        """
        Get summary statistics from parsed data

        Args:
            parsed_data: Parsed presentation data

        Returns:
            Summary dictionary
        """
        if not parsed_data.get('success'):
            return {'error': 'Failed to parse file'}

        slides = parsed_data.get('slides', [])

        return {
            'total_slides': parsed_data.get('total_slides', 0),
            'total_words': parsed_data.get('total_words', 0),
            'total_chars': parsed_data.get('total_chars', 0),
            'avg_words_per_slide': parsed_data.get('total_words', 0) / max(len(slides), 1),
            'file_type': parsed_data.get('file_type', 'unknown'),
            'slides_with_content': sum(1 for s in slides if s.get('word_count', 0) > 0)
        }
